from __future__ import annotations
import os, re, csv, io, json, time, uuid, hashlib, logging, asyncio
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from functools import lru_cache
import httpx
from dotenv import dotenv_values
from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.responses import HTMLResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

ROOT = Path(__file__).parent
# Read the project-local .env at request time; do not copy it into os.environ at import.
ENV_PATH = ROOT / ".env"
DATA = ROOT / "data"
UPLOADS = DATA / "uploads"
LOGS = DATA / "logs"
for p in (UPLOADS, LOGS): p.mkdir(parents=True, exist_ok=True)

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
logger = logging.getLogger("doc-rag")
app = FastAPI(title="Document RAG + RAGAS Evaluation", version="1.0.0")
app.mount("/static", StaticFiles(directory=str(ROOT / "static")), name="static")

# In-memory session index for a simple local demo. Uploaded source files and audit logs persist on disk.
SESSIONS: dict[str, dict[str, Any]] = {}
MAX_FILE_MB = int(os.getenv("MAX_FILE_MB", "25"))
CHUNK_SIZE = int(os.getenv("CHUNK_SIZE", "900"))
CHUNK_OVERLAP = int(os.getenv("CHUNK_OVERLAP", "150"))
DEFAULT_GROQ_MODEL = "openai/gpt-oss-20b"
GROQ_URL = "https://api.groq.com/openai/v1/chat/completions"


def groq_config() -> tuple[str, str, str]:
    """Read config at request time, so editing .env doesn't require a process restart."""
    file_values = dotenv_values(ENV_PATH) if ENV_PATH.exists() else {}
    # A non-empty process environment variable wins; otherwise use the project .env.
    env_key = (os.getenv("GROQ_API_KEY") or "").strip()
    file_key = str(file_values.get("GROQ_API_KEY") or "").strip()
    key = env_key or file_key
    env_model = (os.getenv("GROQ_MODEL") or "").strip()
    file_model = str(file_values.get("GROQ_MODEL") or "").strip()
    model = env_model or file_model or DEFAULT_GROQ_MODEL
    source = "environment" if env_key else (".env" if file_key else "none")
    # Prevent common placeholder values from being treated as configured credentials.
    if key.lower() in {"your_groq_api_key", "your_actual_groq_api_key_here", "gsk_your_actual_key_here", "paste_your_key_here", "changeme", "none", "null"}:
        key = ""
        source = "placeholder/empty"
    return key, model, source


def audit(event: str, **details):
    row = {"timestamp_utc": datetime.now(timezone.utc).isoformat(), "event": event, **details}
    path = LOGS / "audit.jsonl"
    with path.open("a", encoding="utf-8") as f: f.write(json.dumps(row, ensure_ascii=False, default=str) + "\n")
    logger.info("%s %s", event, json.dumps(details, ensure_ascii=False, default=str))

def safe_name(name: str) -> str:
    return Path(name or "uploaded_file").name.replace("\x00", "")[:180]

def extract_file(path: Path, original_name: str) -> tuple[str, list[dict]]:
    ext = Path(original_name).suffix.lower()
    text = ""
    evidence = []
    try:
        if ext == ".pdf":
            import fitz
            doc = fitz.open(path)
            pages = []
            for i, page in enumerate(doc):
                t = page.get_text("text") or ""
                pages.append(t)
                if t.strip(): evidence.append({"source": original_name, "location": f"page {i+1}", "text": t})
            text = "\n\n".join(f"[PAGE {i+1}]\n{t}" for i,t in enumerate(pages))
        elif ext in (".docx",):
            from docx import Document
            d = Document(path)
            parts = []
            for i,p in enumerate(d.paragraphs):
                if p.text.strip(): parts.append(p.text); evidence.append({"source": original_name, "location": f"paragraph {i+1}", "text": p.text})
            for ti,table in enumerate(d.tables, 1):
                for ri,row in enumerate(table.rows, 1):
                    line = " | ".join(c.text.replace("\n"," ") for c in row.cells)
                    parts.append(line); evidence.append({"source": original_name, "location": f"table {ti}, row {ri}", "text": line})
            text = "\n".join(parts)
        elif ext in (".xlsx", ".xlsm"):
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
            parts = []
            for ws in wb.worksheets:
                parts.append(f"[SHEET: {ws.title}]")
                for ri,row in enumerate(ws.iter_rows(values_only=True), 1):
                    vals = ["" if v is None else str(v) for v in row]
                    line = " | ".join(vals)
                    if line.strip(" |"): parts.append(line); evidence.append({"source": original_name, "location": f"sheet {ws.title}, row {ri}", "text": line})
            text = "\n".join(parts)
        elif ext == ".pptx":
            from pptx import Presentation
            prs = Presentation(path); parts = []
            for i,slide in enumerate(prs.slides, 1):
                lines = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip(): lines.append(shape.text.strip())
                    if getattr(shape, "has_table", False):
                        for row in shape.table.rows: lines.append(" | ".join(c.text for c in row.cells))
                line = "\n".join(lines); parts.append(f"[SLIDE {i}]\n{line}")
                if line.strip(): evidence.append({"source": original_name, "location": f"slide {i}", "text": line})
            text = "\n\n".join(parts)
        elif ext == ".csv":
            with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as f:
                rows = list(csv.reader(f))
            lines = [" | ".join(r) for r in rows]
            text = "\n".join(lines)
            evidence = [{"source": original_name, "location": f"row {i+1}", "text": line} for i,line in enumerate(lines) if line.strip()]
        elif ext in (".txt", ".md", ".json", ".xml", ".html", ".yaml", ".yml", ".log", ".ini", ".py", ".sql"):
            text = path.read_text(encoding="utf-8", errors="replace")
            evidence = [{"source": original_name, "location": f"line {i+1}", "text": line} for i,line in enumerate(text.splitlines()) if line.strip()]
        elif ext in (".png", ".jpg", ".jpeg", ".tiff", ".tif", ".bmp", ".webp"):
            try:
                from PIL import Image
                import pytesseract
                text = pytesseract.image_to_string(Image.open(path))
                evidence = [{"source": original_name, "location": "OCR extracted text", "text": text}] if text.strip() else []
            except Exception as e:
                raise ValueError("Image OCR needs Pillow, pytesseract, and the Tesseract OCR system binary installed.") from e
        else:
            raise ValueError(f"Unsupported file type: {ext or '(no extension)'}. Supported: PDF, DOCX, XLSX/XLSM, PPTX, CSV, TXT, MD, JSON, XML, HTML, YAML, LOG, PY, SQL, and common images.")
    except ValueError:
        raise
    except Exception as e:
        raise ValueError(f"Could not extract {original_name}: {type(e).__name__}: {e}") from e
    text = text.strip()
    # Scanned PDFs often contain images but no text layer. Try OCR only when normal extraction is empty.
    if not text and ext == ".pdf":
        try:
            import fitz
            from PIL import Image
            import pytesseract
            doc = fitz.open(path)
            ocr_pages = []
            for i, page in enumerate(doc, 1):
                pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                image = Image.open(io.BytesIO(pix.tobytes("png")))
                page_text = pytesseract.image_to_string(image).strip()
                if page_text:
                    ocr_pages.append(f"[PAGE {i}]\n{page_text}")
                    evidence.append({"source": original_name, "location": f"page {i} (OCR)", "text": page_text})
            text = "\n\n".join(ocr_pages).strip()
        except Exception as exc:
            logger.warning("PDF OCR fallback unavailable: %s", type(exc).__name__)
    if not text:
        raise ValueError("No text extracted. The file may be empty, image-only, encrypted, or unsupported. For scanned PDFs/images install Tesseract OCR and ensure tesseract.exe is on PATH.")
    return text, evidence

def chunk_text(text: str, size=CHUNK_SIZE, overlap=CHUNK_OVERLAP) -> list[str]:
    text = re.sub(r"[ \t]+", " ", text).strip()
    if not text: return []
    out=[]; start=0
    while start < len(text):
        end=min(start+size, len(text))
        if end < len(text):
            split=text.rfind("\n", start, end)
            if split > start + size//2: end=split
        piece=text[start:end].strip()
        if piece: out.append(piece)
        if end >= len(text): break
        start=max(end-overlap, start+1)
    return out

def retrieve(question: str, chunks: list[dict], k=5):
    """Lightweight intent-expanded lexical retrieval; scores are not semantic similarity."""
    stop = {"what", "are", "the", "is", "a", "an", "of", "in", "on", "for", "to", "and", "or", "from", "me", "find", "show", "tell", "give", "please", "document", "this", "that", "with", "about"}
    q = question.lower()
    terms = set(re.findall(r"[a-z0-9]+", q)) - stop
    expansions = {
        "date": {"date", "dates", "validity", "valid", "starts", "start", "expires", "expiry", "expired", "renew", "before", "issued", "endorsed", "june", "january", "february", "march", "april", "may", "july", "august", "september", "october", "november", "december", "2026", "2027", "2028", "2029", "2030", "2031"},
        "dates": {"date", "dates", "validity", "valid", "starts", "start", "expires", "expiry", "renew", "issued", "endorsed"},
        "key": {"policy", "validity", "date", "amount", "name", "number", "details"},
        "amount": {"amount", "value", "premium", "rupees", "rs", "lac", "deductible", "insured", "cost", "price", "85", "118"},
        "amounts": {"amount", "value", "premium", "rupees", "rs", "lac", "deductible", "insured", "cost", "price"},
        "name": {"name", "owner", "insured", "policyholder", "sandeep", "acko", "bike"},
        "names": {"name", "owner", "insured", "policyholder", "sandeep", "acko", "bike"},
        "id": {"id", "number", "policy", "cin", "uin", "irdai", "registration", "gstin", "hsn"},
        "ids": {"id", "number", "policy", "cin", "uin", "irdai", "registration", "gstin", "hsn"},
        "policy": {"policy", "number", "validity", "insured", "owner", "premium", "starts", "expires"},
        "summarize": {"summary", "coverage", "policy", "coverages", "insured", "bike", "owner", "validity"},
        "summary": {"summary", "coverage", "policy", "coverages", "insured", "bike", "owner", "validity"},
    }
    expanded = set(terms)
    for term in list(terms): expanded |= expansions.get(term, set())
    if any(x in q for x in ("date", "when", "expires", "validity", "starts")):
        expanded |= expansions["date"]
    if any(x in q for x in ("amount", "cost", "price", "premium", "value")):
        expanded |= expansions["amount"]
    if any(x in q for x in ("name", "owner", "who", "id", "number", "policy")):
        expanded |= expansions["id"] | expansions["name"]
    scored=[]
    total_docs=max(1,len(chunks))
    for c in chunks:
        text=c["text"].lower()
        ct=set(re.findall(r"[a-z0-9]+", text))
        overlap=expanded & ct
        if not overlap: continue
        # Give original question terms more weight than expansion-only terms.
        direct=len(terms & ct)
        coverage=len(overlap)/max(1,len(expanded))
        score=(direct*2.0 + len(overlap)*0.65)/max(1, len(terms)**0.5)
        if re.search(r"\b(19|20)\d{2}\b|\b\d{1,2}(st|nd|rd|th)?\s+(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)", text, re.I) and expanded & expansions["date"]:
            score += 1.0
        if re.search(r"\b(policy number|cin|uin|gstin|irdai reg no|insured value|premium|bike owner|name)\b", text, re.I) and expanded & (expansions["id"]|expansions["name"]|expansions["amount"]):
            score += 0.7
        scored.append((score, c))
    scored.sort(key=lambda x:x[0], reverse=True)
    return [{"score": round(s,4), **c} for s,c in scored[:k]]

async def call_groq(messages: list[dict], timeout: float = 45) -> dict:
    key, model, source = groq_config()
    if not key:
        raise HTTPException(status_code=503, detail=(
            "Groq API key is not loaded. Add GROQ_API_KEY=gsk_... to the .env file located beside app.py, "
            "save it, then retry. Ensure the file is named exactly .env (not .env.txt)."
        ))
    try:
        async with httpx.AsyncClient(timeout=timeout) as client:
            response = await client.post(
                GROQ_URL,
                headers={"Authorization": f"Bearer {key}", "Content-Type": "application/json"},
                json={"model": model, "messages": messages, "temperature": 0.1},
            )
    except httpx.TimeoutException:
        raise HTTPException(status_code=504, detail="Groq connection timed out. Check your internet connection and retry.")
    except httpx.RequestError as exc:
        raise HTTPException(status_code=502, detail=f"Could not connect to api.groq.com: {type(exc).__name__}. Check internet, proxy, or firewall.")
    if response.status_code >= 400:
        try:
            payload = response.json()
            api_message = payload.get("error", {}).get("message", response.text[:500])
            api_code = payload.get("error", {}).get("code", "")
        except Exception:
            api_message, api_code = response.text[:500], ""
        audit("llm_error", status_code=response.status_code, api_code=api_code, message=str(api_message)[:500], model=model, key_source=source)
        if response.status_code in (401, 403):
            detail = "Groq rejected the API key (401/403). Create a fresh key in https://console.groq.com/keys and check that .env contains the key without extra text."
        elif response.status_code == 429:
            detail = "Groq rate limit/quota reached (429). Wait, check your Groq limits, then retry. API says: " + str(api_message)[:300]
        elif response.status_code == 400 and ("model" in str(api_message).lower() or "model" in str(api_code).lower()):
            detail = f"Groq model is not accepted: {model}. Change GROQ_MODEL in .env to a model enabled for your Groq account. API says: {str(api_message)[:300]}"
        else:
            detail = f"Groq API returned HTTP {response.status_code}: {str(api_message)[:400]}"
        raise HTTPException(status_code=502, detail=detail)
    try:
        return response.json()
    except ValueError:
        raise HTTPException(status_code=502, detail="Groq returned a non-JSON response.")


async def groq_answer(question: str, context: list[dict]) -> str:
    if not context:
        return "I couldn't find relevant text in the uploaded documents. Try a more specific question or verify that the file's text was extracted."
    key, _, _ = groq_config()
    if not key:
        return "Groq is not configured, so this is retrieval-only output (not an LLM-generated answer). Relevant source excerpts:\n\n" + "\n\n".join(
            f"[{c['source']} — {c['location']}]\n{c['text']}" for c in context[:5]
        )
    joined = "\n\n".join(
        f"[{i+1}] Source: {c['source']} | Location: {c['location']} | Chunk: {c['chunk_id']}\n{c['text']}"
        for i, c in enumerate(context)
    )
    prompt = (
        "Answer only using the supplied document excerpts. If the answer is absent, say it is not found. "
        "Do not guess. Cite each factual claim with source and location like [filename, page 2]. "
        "If the document has contradictory or malformed text, explicitly flag the uncertainty.\n\n"
        f"DOCUMENT EXCERPTS:\n{joined}\n\nQUESTION: {question}"
    )
    data = await call_groq([
        {"role": "system", "content": "You are a careful document QA assistant. Ground every claim in the provided context."},
        {"role": "user", "content": prompt},
    ])
    try:
        return data["choices"][0]["message"]["content"]
    except (KeyError, IndexError, TypeError):
        raise HTTPException(status_code=502, detail="Groq response did not contain an answer. Check model compatibility.")

@lru_cache(maxsize=4)
def build_ragas_evaluators(api_key: str, model: str):
    """Build genuine RAGAS metrics; heavy embedding model is loaded only once per process."""
    from langchain_groq import ChatGroq
    from langchain_huggingface import HuggingFaceEmbeddings
    from ragas.embeddings import LangchainEmbeddingsWrapper
    from ragas.llms import LangchainLLMWrapper
    from ragas.metrics import Faithfulness, LLMContextPrecisionWithoutReference
    try:
        from ragas.metrics import AnswerRelevancy
    except ImportError:
        from ragas.metrics import ResponseRelevancy as AnswerRelevancy

    judge = ChatGroq(model=model, api_key=api_key, temperature=0, max_tokens=1024)
    ragas_llm = LangchainLLMWrapper(judge)
    local_embeddings = HuggingFaceEmbeddings(
        model_name="sentence-transformers/all-MiniLM-L6-v2",
        model_kwargs={"device": "cpu"},
        encode_kwargs={"normalize_embeddings": True},
    )
    ragas_embeddings = LangchainEmbeddingsWrapper(local_embeddings)
    return {
        "faithfulness": Faithfulness(llm=ragas_llm),
        "answer_relevancy": AnswerRelevancy(llm=ragas_llm, embeddings=ragas_embeddings),
        "context_precision": LLMContextPrecisionWithoutReference(llm=ragas_llm),
    }

async def evaluate_evidence_coverage(question: str, answer: str, context: list[dict]) -> dict:
    """Custom evidence audit, explicitly separate from official RAGAS metrics."""
    evidence = "\n\n".join(
        f"[{i+1}] {c['source']} — {c['location']}\n{c['text']}" for i, c in enumerate(context[:6])
    )
    prompt = f"""Audit whether the answer's important factual details have direct support in the supplied evidence.
Do not use outside knowledge. Count factual claims in the answer, then count how many are directly supported by the evidence. A detail that is merely plausible is not supported. If the answer says it cannot be found, treat that as appropriate only if the evidence truly does not contain the requested detail.
Return only JSON: {{"total_claims": 0, "supported_claims": 0, "unsupported_claims": ["exact or concise claim"], "issues": ["reason"]}}.
QUESTION: {question}\nANSWER: {answer}\nEVIDENCE:\n{evidence}"""
    data = await call_groq([
        {"role": "system", "content": "You are a conservative evidence auditor. Never invent claims or scores."},
        {"role": "user", "content": prompt},
    ], timeout=45)
    raw = data["choices"][0]["message"]["content"]
    match = re.search(r"\{.*\}", raw, re.S)
    if not match:
        raise ValueError("Evidence auditor returned no JSON object")
    parsed = json.loads(match.group(0))
    total = int(parsed.get("total_claims", 0))
    supported = int(parsed.get("supported_claims", 0))
    if total < 0 or supported < 0 or supported > total:
        raise ValueError("Evidence auditor returned invalid claim counts")
    # No factual claims is not automatically 100%; mark score unavailable to avoid a misleading result.
    coverage = round(100 * supported / total, 1) if total > 0 else None
    return {
        "score": coverage,
        "total_claims": total,
        "supported_claims": supported,
        "unsupported_claims": parsed.get("unsupported_claims", [])[:12],
        "issues": parsed.get("issues", [])[:12],
    }

async def evaluate_answer_live(question: str, answer: str, context: list[dict]) -> dict:
    """Run official RAGAS metrics after each answer; never substitute guessed scores on failure."""
    if not context:
        return {"status": "not_evaluated", "engine": "RAGAS", "metrics": {}, "note": "No retrieved evidence was available, so metrics were not calculated."}
    key, model, _ = groq_config()
    if not key:
        return {"status": "unavailable", "engine": "RAGAS", "metrics": {}, "note": "Groq key unavailable; no evaluation scores were generated."}

    try:
        from ragas import SingleTurnSample
        evaluators = await asyncio.to_thread(build_ragas_evaluators, key, model)
        sample = SingleTurnSample(
            user_input=question,
            response=answer,
            retrieved_contexts=[c["text"] for c in context[:6]],
        )
        metrics = {"faithfulness": None, "answer_relevancy": None, "context_precision": None, "evidence_coverage": None}
        metric_errors = {}
        # Each score comes from the actual RAGAS metric implementation, not a hand-written replacement.
        for name, metric in evaluators.items():
            try:
                value = await asyncio.wait_for(metric.single_turn_ascore(sample), timeout=90)
                if hasattr(value, "value"):
                    value = value.value
                value = float(value)
                if not (value == value and abs(value) != float("inf")):
                    raise ValueError("RAGAS returned a non-finite score")
                # Some relevancy implementations can mathematically return a value outside 0..1;
                # preserve its actual value rather than silently clipping it.
                metrics[name] = round(value * 100, 1)
            except Exception as exc:
                logger.exception("RAGAS metric failed: %s", name)
                metric_errors[name] = f"{type(exc).__name__}: {str(exc)[:220]}"
        unsupported_claims=[]
        issues=[]
        evidence_claim_counts=None
        try:
            coverage = await evaluate_evidence_coverage(question, answer, context)
            metrics["evidence_coverage"] = coverage["score"]
            unsupported_claims = coverage["unsupported_claims"]
            issues = coverage["issues"]
            evidence_claim_counts = {"total_claims": coverage["total_claims"], "supported_claims": coverage["supported_claims"]}
        except Exception as exc:
            logger.exception("Custom evidence coverage audit failed")
            metric_errors["evidence_coverage"] = f"{type(exc).__name__}: {str(exc)[:220]}"
        scored = [v for v in metrics.values() if v is not None]
        if len(scored) == len(metrics):
            status = "completed"
        elif scored:
            status = "partial"
        else:
            status = "failed"
        # Don't manufacture a verdict. The user can inspect the source and each score's provenance.
        verdict = None
        if metrics["faithfulness"] is not None:
            verdict = "high_support" if metrics["faithfulness"] >= 85 else ("mixed_support" if metrics["faithfulness"] >= 60 else "low_support")
        return {
            "status": status,
            "engine": "Official RAGAS Python metrics + separate custom evidence audit",
            "method": {
                "faithfulness": "RAGAS Faithfulness",
                "answer_relevancy": "RAGAS AnswerRelevancy / ResponseRelevancy",
                "context_precision": "RAGAS LLMContextPrecisionWithoutReference",
                "evidence_coverage": "Custom claim-count audit; not an official RAGAS metric",
            },
            "metrics": metrics,
            "metric_errors": metric_errors,
            "evidence_claim_counts": evidence_claim_counts,
            "unsupported_claims": unsupported_claims,
            "issues": issues,
            "verdict": verdict,
            "note": "Faithfulness, answer relevancy, and context precision are calculated with the installed RAGAS Python library. Evidence coverage is a separate LLM-assisted claim audit, not an official RAGAS metric. These scores are estimates against retrieved text, not proof that the source document itself is true. Reference-based correctness/recall requires trusted reference answers.",
        }
    except Exception as exc:
        logger.exception("RAGAS initialization/evaluation failed")
        return {
            "status": "failed", "engine": "RAGAS", "metrics": {},
            "note": f"No RAGAS scores were generated. Check dependencies/model download/Groq limits. Error: {type(exc).__name__}: {str(exc)[:350]}",
        }

class ChatRequest(BaseModel):
    session_id: str
    question: str

@app.get("/", response_class=HTMLResponse)
async def home():
    return (ROOT / "static" / "index.html").read_text(encoding="utf-8")

@app.get("/api/health")
async def health():
    key, model, source = groq_config()
    return {
        "status": "ok",
        "groq_configured": bool(key),
        "groq_key_source": source,
        "env_file_exists": ENV_PATH.exists(),
        "env_file_path": str(ENV_PATH),
        "model": model if key else None,
        "message": "A non-placeholder key was detected. Run the Groq connection test to verify credentials, model access, and network connectivity." if key else "No usable key detected. Check the .env file beside app.py; its filename must be exactly .env (not .env.txt).",
    }


@app.post("/api/groq/test")
async def test_groq_connection():
    key, model, source = groq_config()
    if not key:
        raise HTTPException(status_code=503, detail="No API key detected. Add GROQ_API_KEY=gsk_... to the .env beside app.py. Check that it is not named .env.txt.")
    data = await call_groq([{"role": "user", "content": "Reply with exactly: GROQ_CONNECTION_OK"}], timeout=25)
    answer = data.get("choices", [{}])[0].get("message", {}).get("content", "")
    audit("groq_connection_test", success=True, model=model, key_source=source)
    return {"connected": True, "model": model, "key_source": source, "message": "Groq API connection successful.", "test_reply": answer[:100]}

@app.post("/api/upload")
async def upload(files: list[UploadFile] = File(...)):
    session_id=str(uuid.uuid4())
    all_chunks=[]; file_info=[]; all_evidence=[]
    for up in files:
        name=safe_name(up.filename)
        content=await up.read()
        if not content: raise HTTPException(400, f"{name}: empty file")
        if len(content) > MAX_FILE_MB*1024*1024: raise HTTPException(413, f"{name}: exceeds {MAX_FILE_MB} MB limit")
        digest=hashlib.sha256(content).hexdigest()
        stored=UPLOADS / f"{uuid.uuid4().hex}_{name}"
        stored.write_bytes(content)
        try:
            text, evidence=extract_file(stored, name)
        except ValueError as e:
            stored.unlink(missing_ok=True)
            audit("upload_failed", filename=name, sha256=digest, error=str(e))
            raise HTTPException(415, detail=str(e))
        chunks=[]
        # Chunk each source block independently so a chunk never accidentally spans PDF pages,
        # spreadsheet rows, slides, or paragraphs with misleading provenance.
        for ev in evidence:
            for piece_index, piece in enumerate(chunk_text(ev["text"]), 1):
                chunks.append({"chunk_id":f"{digest[:10]}-{len(chunks)+1}","source":name,
                               "location":ev["location"],"block_piece":piece_index,"text":piece})
        # Some parsers may produce extracted text without evidence blocks; keep it retrievable.
        if not chunks and text.strip():
            for piece in chunk_text(text):
                chunks.append({"chunk_id":f"{digest[:10]}-{len(chunks)+1}","source":name,"location":"document text","text":piece})
        file_info.append({"filename":name,"bytes":len(content),"sha256":digest,"characters_extracted":len(text),"chunks":len(chunks),"evidence_items":len(evidence),"stored_file":stored.name})
        all_chunks.extend(chunks)
        all_evidence.extend(evidence)
        audit("upload_success", session_id=session_id, filename=name, sha256=digest, bytes=len(content), chars=len(text), chunks=len(chunks), evidence_items=len(evidence))
    SESSIONS[session_id]={"files":file_info,"chunks":all_chunks,"evidence":all_evidence,"created_at":datetime.now(timezone.utc).isoformat(),"history":[]}
    return {"session_id":session_id,"files":file_info,"chunk_count":len(all_chunks),"evidence_count":len(all_evidence),"message":"Files extracted and indexed for this running session."}

@app.post("/api/chat")
async def chat(req: ChatRequest):
    session=SESSIONS.get(req.session_id)
    if not session: raise HTTPException(404, "Session not found. Upload the documents again after restarting the server.")
    q=req.question.strip()
    if not q: raise HTTPException(400,"Question cannot be empty")
    t0=time.perf_counter()
    found=retrieve(q,session["chunks"])
    context=[{"source":x["source"],"location":x["location"],"text":x["text"],"score":x["score"],"chunk_id":x["chunk_id"]} for x in found]
    answer=await groq_answer(q,context)
    # Evaluate immediately so users see the evidence-based quality estimate with every answer.
    evaluation = await evaluate_answer_live(q, answer, context) if groq_config()[0] else {"status":"unavailable","engine":"none","note":"Configure a working Groq key to generate answers and run automatic evaluation."}
    _, active_model, _ = groq_config()
    record={"request_id":str(uuid.uuid4()),"timestamp_utc":datetime.now(timezone.utc).isoformat(),"question":q,"answer":answer,
            "retrieved_context":context,"retrieval_count":len(context),"latency_ms":round((time.perf_counter()-t0)*1000,2),
            "model":active_model,"evaluation":evaluation}
    session["history"].append(record)
    audit("chat_answer",session_id=req.session_id,**record)
    return record

@app.get("/api/session/{session_id}/logs")
async def session_logs(session_id: str):
    session=SESSIONS.get(session_id)
    if not session: raise HTTPException(404,"Session not found")
    return {"session_id":session_id,"files":session["files"],"evidence":session["evidence"],"history":session["history"]}

@app.post("/api/session/{session_id}/evaluate")
async def evaluate(session_id: str):
    session=SESSIONS.get(session_id)
    if not session: raise HTTPException(404,"Session not found")
    if not session["history"]: raise HTTPException(400,"Ask at least one question before evaluation.")
    completed=[h for h in session["history"] if h.get("evaluation",{}).get("status") in {"completed", "partial"}]
    metric_names=("faithfulness","answer_relevancy","context_precision","evidence_coverage")
    aggregate={}
    for name in metric_names:
        values=[h["evaluation"]["metrics"][name] for h in completed if isinstance(h.get("evaluation",{}).get("metrics",{}).get(name), (int, float))]
        aggregate[name]={"mean_percent":round(sum(values)/len(values),1) if values else None,"scored_answers":len(values)}
    result={"session_id":session_id,"evaluated_at_utc":datetime.now(timezone.utc).isoformat(),
            "status":"completed_estimates" if completed else "evaluation_unavailable",
            "engine":"Official RAGAS Python metrics plus separate custom evidence-coverage audit",
            "note":"Faithfulness, answer relevancy, and context precision use the RAGAS Python library. Evidence coverage is a separate custom claim-support audit and is not an official RAGAS metric. These scores estimate support against retrieved text; they do not prove the source document itself is true. Reference-based context recall and answer correctness require trusted reference answers / labeled data.",
            "questions_evaluated":len(session["history"]),"questions_with_scores":len(completed),
            "aggregate_metrics":aggregate,
            "reference_based_metrics":{"context_recall":"requires trusted reference answer","context_precision_with_reference":"requires trusted reference answer","answer_correctness":"requires trusted reference answer","factual_correctness":"requires trusted reference answer or verified labels"},
            "per_question":[{"request_id":h["request_id"],"question":h["question"],"answer":h["answer"],"retrieved_context":h["retrieved_context"],"retrieval_count":h["retrieval_count"],"evaluation":h.get("evaluation",{"status":"not_available"})} for h in session["history"]]}
    audit("evaluation_report",session_id=session_id,questions=len(session["history"]),scored=len(completed),status=result["status"])
    report_path=LOGS / f"evaluation_{session_id}.json"
    report_path.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding="utf-8")
    return result

@app.get("/api/logs/download")
async def download_logs():
    path=LOGS/"audit.jsonl"
    if not path.exists(): path.write_text("",encoding="utf-8")
    from fastapi.responses import FileResponse
    return FileResponse(path, media_type="application/x-ndjson", filename="audit.jsonl")

@app.get("/api/session/{session_id}/report")
async def download_report(session_id: str):
    path=LOGS/f"evaluation_{session_id}.json"
    if not path.exists(): raise HTTPException(404,"Run evaluation first")
    from fastapi.responses import FileResponse
    return FileResponse(path, media_type="application/json", filename=f"evaluation_{session_id}.json")
